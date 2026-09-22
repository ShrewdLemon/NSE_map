"""Emit the classified register in the Book1 layout.

Book1 columns: Company name | Share holder name | Share holder category |
then (No of shares, Value) per quarter Q1..Q6. Value is intentionally left
blank - the Bloomberg export carries share counts only.
"""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
QTR_FILL = PatternFill("solid", fgColor="2E5496")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def write(rows, quarter_labels, company, path, include_audit=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shareholder Categories"

    base = ["Company name", "Share holder name", "Share holder category"]
    ws.append([])  # row 1 holds the quarter group labels
    ws.append(base + [c for q in quarter_labels for c in ("No of shares", "Value")])

    for i, q in enumerate(quarter_labels):
        col = len(base) + 1 + i * 2
        cell = ws.cell(row=1, column=col, value=q)
        cell.font, cell.fill = WHITE_BOLD, QTR_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)

    for cell in ws[2]:
        if cell.value:
            cell.font, cell.fill = WHITE_BOLD, HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for r in rows:
        line = [company, r["holder_name"], r["category"]]
        for q in quarter_labels:
            line += [r["quarters"].get(q), None]  # Value deliberately blank
        ws.append(line)

    if include_audit:
        start = len(base) + 1 + len(quarter_labels) * 2
        for j, h in enumerate(("Basis", "Confidence", "Country", "Entity type",
                               "Holding vehicle", "Review?", "Source")):
            c = ws.cell(row=2, column=start + j, value=h)
            c.font, c.fill = WHITE_BOLD, HEADER_FILL
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, r in enumerate(rows):
            vals = (r.get("basis"), r.get("confidence"), r.get("country"),
                    r.get("entity_type"), r.get("holding_vehicle"),
                    "REVIEW" if r.get("needs_review") else "", r.get("source"))
            for j, v in enumerate(vals):
                c = ws.cell(row=3 + i, column=start + j, value=v)
                if r.get("needs_review"):
                    c.fill = FLAG_FILL

    widths = {1: 24, 2: 46, 3: 22}
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 15)
    ws.freeze_panes = "D3"
    for row in ws.iter_rows(min_row=3, min_col=4, max_col=len(base) + len(quarter_labels) * 2):
        for c in row:
            c.number_format = "#,##0"

    wb.save(path)
    return path
