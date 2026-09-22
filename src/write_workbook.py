"""Emit every classified company into one workbook.

Sheets:
    Summary       one row per company, category counts across the taxonomy
    Index totals  the whole index rolled up by category
    Provenance    per company: source, as-of date, share base and how it settled
    All holders   every row, Book1 layout with a ticker column in front
    <TICKER>      one sheet per company, exactly the single-company layout
"""
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
QTR_FILL = PatternFill("solid", fgColor="2E5496")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)

AUDIT = ("Basis", "Confidence", "Country", "Entity type", "Holding vehicle",
         "Review?", "Source")
_ILLEGAL = re.compile(r"[\[\]:*?/\\]")


def sheet_name(ticker):
    return _ILLEGAL.sub("-", ticker)[:31]


def _style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        if cell.value is not None:
            cell.font, cell.fill = WHITE_BOLD, HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _holder_block(ws, rows, quarters, lead_cols, lead_values):
    """Write the Book1 holder table. lead_cols are the columns before the quarters."""
    base = list(lead_cols) + ["Share holder name", "Share holder category"]
    ws.append([])  # row 1 carries the merged quarter labels
    ws.append(base + [c for q in quarters for c in ("No of shares", "Value")]
              + list(AUDIT))

    for i, q in enumerate(quarters):
        col = len(base) + 1 + i * 2
        cell = ws.cell(row=1, column=col, value=q)
        cell.font, cell.fill = WHITE_BOLD, QTR_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)

    _style_header(ws, 2, len(base) + len(quarters) * 2 + len(AUDIT))

    for r in rows:
        line = list(lead_values(r)) + [r["holder_name"], r["category"]]
        for q in quarters:
            line += [r["quarters"].get(q), None]  # Value deliberately blank
        line += [r.get("basis"), r.get("confidence"), r.get("country"),
                 r.get("entity_type"), r.get("holding_vehicle"),
                 "REVIEW" if r.get("needs_review") else "", r.get("source")]
        ws.append(line)
        if r.get("needs_review"):
            for c in range(len(base) + len(quarters) * 2 + 1,
                           len(base) + len(quarters) * 2 + len(AUDIT) + 1):
                ws.cell(row=ws.max_row, column=c).fill = FLAG_FILL

    first_qtr = len(base) + 1
    for row in ws.iter_rows(min_row=3, min_col=first_qtr,
                            max_col=len(base) + len(quarters) * 2):
        for c in row:
            c.number_format = "#,##0"
    widths = {len(lead_cols) + 1: 46, len(lead_cols) + 2: 22}
    for i, _ in enumerate(lead_cols, start=1):
        widths[i] = 30 if i == len(lead_cols) else 14
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 15)
    ws.freeze_panes = ws.cell(row=3, column=len(base) + 1).coordinate
    ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"


def write(companies, categories, path, index_name="NIFTY 50"):
    """companies: list of dicts with ticker/company/sector/quarters/rows/meta."""
    quarters_all = []
    for c in companies:
        for q in c["quarters"]:
            if q not in quarters_all:
                quarters_all.append(q)

    wb = openpyxl.Workbook()

    # --- Summary -----------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    head = (["Ticker", "Company name", "Sector", "As at", "Holders named",
             "Shares (SCRR base)", "Named shares", "Named cover %"]
            + list(categories) + ["Unclassified", "Needs review"])
    ws.append(head)
    _style_header(ws, 1, len(head))
    for c in companies:
        counts = {k: 0 for k in categories}
        unclassified = 0
        for r in c["rows"]:
            if r["category"] in counts:
                counts[r["category"]] += 1
            else:
                unclassified += 1
        latest = c["quarters"][-1]
        named = sum(v for r in c["rows"] for v in [r["quarters"].get(latest)] if v)
        base = c["meta"].get("shares_scrr")
        ws.append([c["ticker"], c["company"], c.get("sector"), c["meta"].get("as_of"),
                   len(c["rows"]), base, named,
                   (named / base if base else None)]
                  + [counts[k] for k in categories]
                  + [unclassified,
                     sum(1 for r in c["rows"] if r.get("needs_review"))])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = "#,##0"
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.number_format = "0.0%"
    for col, w in {1: 13, 2: 42, 3: 24, 4: 12, 5: 14, 6: 20, 7: 18, 8: 13}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for col in range(9, len(head) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ws.max_row}"

    # --- Index totals ------------------------------------------------------
    ws = wb.create_sheet("Index totals")
    ws.append([f"{index_name} ownership rolled up by category"])
    ws["A1"].font = BOLD
    ws.append([])
    head = ["Share holder category", "Holder rows", "Companies present",
            "Shares held (latest quarter)", "% of all named shares"]
    ws.append(head)
    _style_header(ws, 3, len(head))

    tally = {k: {"rows": 0, "cos": set(), "shares": 0} for k in categories}
    tally["(unclassified)"] = {"rows": 0, "cos": set(), "shares": 0}
    for c in companies:
        latest = c["quarters"][-1]
        for r in c["rows"]:
            key = r["category"] if r["category"] in tally else "(unclassified)"
            tally[key]["rows"] += 1
            tally[key]["cos"].add(c["ticker"])
            tally[key]["shares"] += r["quarters"].get(latest) or 0
    grand = sum(v["shares"] for v in tally.values()) or 1
    for k, v in sorted(tally.items(), key=lambda kv: -kv[1]["shares"]):
        if not v["rows"]:
            continue
        ws.append([k, v["rows"], len(v["cos"]), v["shares"], v["shares"] / grand])
    ws.append(["Total", sum(v["rows"] for v in tally.values()), len(companies),
               grand, 1.0])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = BOLD, TOTAL_FILL
    for row in ws.iter_rows(min_row=4, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"
    for row in ws.iter_rows(min_row=4, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0.00%"
    for col, w in {1: 26, 2: 13, 3: 18, 4: 26, 5: 20}.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    # --- Provenance --------------------------------------------------------
    # Fifty companies read out of filings by machine need their sources on the
    # face of the workbook, not buried in a repository.
    ws = wb.create_sheet("Provenance")
    head = ["Ticker", "Company name", "ISIN", "As at", "Quarters",
            "Shares (SCRR base)", "How the share base was settled",
            "Promoter % filed", "Holders named", "Source"]
    ws.append(head)
    _style_header(ws, 1, len(head))
    for c in companies:
        m = c["meta"]
        ws.append([c["ticker"], c["company"], m.get("isin"), m.get("as_of"),
                   ", ".join(c["quarters"]), m.get("shares_scrr"),
                   m.get("denominator_note"), m.get("promoter_pct"),
                   len(c["rows"]), "\n".join(m.get("source_urls") or [])])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=10):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in {1: 13, 2: 40, 3: 15, 4: 12, 5: 34, 6: 20, 7: 52,
                   8: 15, 9: 14, 10: 60}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{ws.max_row}"

    # --- All holders -------------------------------------------------------
    ws = wb.create_sheet("All holders")
    every = [r for c in companies for r in c["rows"]]
    lookup = {id(r): c for c in companies for r in c["rows"]}
    _holder_block(ws, every, quarters_all, ["Ticker", "Company name"],
                  lambda r: (lookup[id(r)]["ticker"], lookup[id(r)]["company"]))

    # --- One sheet per company --------------------------------------------
    for c in companies:
        cs = wb.create_sheet(sheet_name(c["ticker"]))
        _holder_block(cs, c["rows"], c["quarters"], ["Company name"],
                      lambda r, co=c: (co["company"],))

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
